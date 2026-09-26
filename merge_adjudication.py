# -*- coding: utf-8 -*-
"""
课题④ 儿童胃十二指肠穿孔 —— 合并双人裁定，计算 Cohen κ，生成仲裁清单与终裁数据集

输入：build_adjudication_sheet.py 生成、两名裁定者各自填好的两份盲表
      裁定工作表_v1_裁定者A.xlsx / 裁定工作表_v1_裁定者B.xlsx（只读，本脚本绝不改写）。
输出：裁定合并_一致性与仲裁_v1.xlsx
  · 说明          来源文件 SHA256、进度与方法说明（可直接改写进论文方法部分）；
  · 逻辑核查      每名裁定者自身填写矛盾之处，应退回本人修改后再合并；
  · 一致性κ       各字段的观察一致率、Cohen κ 及 95% CI；
  · 病因大类交叉表 A 行 × B 列，看哪些类别最易混淆，用于预试后校准分类标准；
  · 分歧与仲裁    逐条列出两人不一致的字段，仲裁者在黄色列填写；
  · 终裁数据集    一致者取共同值，不一致者取仲裁结果，尚未仲裁者标【待仲裁】。

可随时运行：预试阶段只填了 10–15 例也能算 κ，只统计两人都已完成的病例。

安全约定（同 build_adjudication_sheet.py）：
  1. 输出文件若已存在，先备份为 *.backup_时间戳.xlsx；
  2. 按（科研就诊编号, 字段）把旧输出里已填写的仲裁结果回填到新表；若回填时发现
     该条的 A/B 值已与上次不同，旧仲裁作废并在“回填提示”中保留原文，待重新仲裁，不静默沿用；
  3. 所有写盘动作只在 __main__ 中触发，import 本模块不产生任何副作用。
"""
import hashlib
import math
import os
import shutil
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, Alignment, PatternFill

import build_adjudication_sheet as adj

KEY = adj.KEY
OUT = r'D:\胃十二指肠穿孔\裁定合并_一致性与仲裁_v1.xlsx'



def _both(col, test):
    """病例条件：两名裁定者在 col 上的取值都满足 test。"""
    return lambda a, b: a[col].map(test).astype(bool) & b[col].map(test).astype(bool)


_inc = _both('是否排除', lambda v: v == '否')
# 需一致性评价与仲裁的字段：(字段, 纳入 κ 计算的病例条件, 条件说明)
# 条件字段只在两人都认为该字段适用的病例里比较，否则大量“NA 对 NA”的一致会把 κ 抬高。
CAT_FIELDS = [
    ('病因大类', None, '两人均已完成的全部病例'),
    ('是否排除', None, '两人均已完成的全部病例'),
    ('排除理由', _both('是否排除', lambda v: v == '是'), '两人均判为排除的病例'),
    ('异物类型', _both('病因大类', lambda v: v.startswith('1')), '两人均判为异物相关的病例'),
    ('穿孔部位', _inc, '两人均判为纳入的病例'),
    ('手术入路', _inc, '两人均判为纳入的病例'),
    ('HP状态', _inc, '两人均判为纳入的病例'),
    ('结局', _inc, '两人均判为纳入的病例'),
]
# 数值字段：只报观察一致率，不算 κ
NUM_FIELDS = [
    ('磁性异物枚数', _both('异物类型', lambda v: v == '磁性'), '两人均判为磁性异物的病例'),
    ('穿孔数目', _inc, '两人均判为纳入的病例'),
]
CAT_NAMES = {f for f, _, _ in CAT_FIELDS}
ARB_FIELDS = [f for f, _, _ in CAT_FIELDS + NUM_FIELDS]
# 自由文本：不仲裁，两人不同时并列保留，供仲裁者或作者整理
TEXT_FIELDS = ['异物具体名称', '合并其他消化道穿孔', '备注']
# 从盲表带入终裁数据集的病例基本信息
INFO_COLS = ['科研患者编号', KEY, '住院号', '第几次住院', '数据完整性标记', '性别', '年龄_岁', '年龄_天',
             '年龄组', '入院日期', '入院年', '住院天数', '入院科别']
ARB_COLS = ['仲裁结果', '仲裁者', '仲裁理由']
PENDING = '【待仲裁】'


def sha256(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for blk in iter(lambda: f.read(1 << 20), b''):
            h.update(blk)
    return h.hexdigest()


def norm(v):
    """统一单元格取值：空值→''，2.0→'2'，去首尾空格，NR 统一大写。"""
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    if s in ('nan', 'NaN', 'None'):
        return ''
    return 'NR' if s.upper() == 'NR' else s


def load_rater(path):
    # keep_default_na=False：否则下拉选项“NA”会被 pandas 当成缺失值读成空
    df = pd.read_excel(path, sheet_name='裁定表', dtype=object, keep_default_na=False)
    miss = [c for c in [KEY] + adj.ADJ_COLS if c not in df.columns]
    assert not miss, '%s 缺少列：%s' % (os.path.basename(path), miss)
    dup = df[KEY][df[KEY].duplicated()]
    assert dup.empty, '%s 就诊编号重复：%s' % (os.path.basename(path), list(dup))
    for c in adj.ADJ_COLS:
        df[c] = df[c].map(norm)
    df[KEY] = df[KEY].map(norm)
    return df.set_index(KEY, drop=False)


# ---------------------------------------------------------------- 逻辑核查
def _pos_int_or_nr(v):
    return v == 'NR' or (v.isdigit() and int(v) > 0)


def logic_check(df, rater):
    """返回该裁定者自身的填写矛盾清单（级别：错误 / 提示）。"""
    out = []

    def add(r, level, msg):
        out.append({'裁定者': rater, '级别': level, KEY: r[KEY], '住院号': r.get('住院号', ''),
                    '入院日期': r.get('入院日期', ''), '问题': msg})

    for _, r in df.iterrows():
        filled = [c for c in adj.ADJ_COLS if r[c] and c not in ('裁定者', '裁定日期')]
        if not filled:
            continue
        cause, excl, reason = r['病因大类'], r['是否排除'], r['排除理由']
        for c, opts in adj.OPTIONS.items():
            if r[c] and r[c] not in opts:
                add(r, '错误', '“%s”填了下拉选项以外的值：%s' % (c, r[c]))
        if bool(cause) != bool(excl):
            add(r, '错误', '病因大类与是否排除须同时填写')
        if cause and excl and (cause.startswith('9') != (excl == '是')):
            add(r, '错误', '病因大类“%s”与是否排除“%s”矛盾' % (cause, excl))
        if excl == '是' and not reason:
            add(r, '错误', '判为排除但未选排除理由')
        if excl == '否' and reason:
            add(r, '错误', '判为纳入却填了排除理由')
        if cause.startswith('1') and r['异物类型'] not in ('磁性', '非磁性'):
            add(r, '错误', '异物相关病例须选磁性或非磁性')
        if cause and not cause.startswith('1') and r['异物类型'] in ('磁性', '非磁性'):
            add(r, '错误', '非异物病例的异物类型应填 NA')
        if r['异物类型'] == '磁性' and not r['磁性异物枚数']:
            add(r, '错误', '磁性异物须填枚数（记不清写 NR）')
        for c in ('磁性异物枚数', '穿孔数目'):
            if r[c] and not _pos_int_or_nr(r[c]):
                add(r, '错误', '“%s”应为正整数或 NR：%s' % (c, r[c]))
        if excl == '否':
            lack = [c for c in ('穿孔部位', '手术入路', '结局') if not r[c]]
            if lack:
                add(r, '错误', '纳入病例缺必填项：%s' % '、'.join(lack))
        if cause.startswith('2'):
            try:
                if float(r['年龄_天']) > 28:
                    add(r, '错误', '判为新生儿自发性，但入院日龄 %s 天＞28 天' % norm(r['年龄_天']))
            except (TypeError, ValueError):
                pass
        if reason.startswith('E5') and ('无出院小结' in str(r.get('数据完整性标记', ''))
                                        or '无手术记录' in str(r.get('数据完整性标记', ''))):
            add(r, '提示', '按 E5 排除且属科研平台缺记录病例：请确认已回 HIS 调阅原始病历')
    return out


# ---------------------------------------------------------------- Cohen κ
def cohen_kappa(x, y):
    """未加权 Cohen κ；95% CI 用 Fleiss, Cohen & Everitt (1969) 大样本标准误。
    返回 (n, 一致数, 观察一致率, κ, CI下限, CI上限)；κ 不可估计时为 NaN。"""
    x, y = list(x), list(y)
    n = len(x)
    if n == 0:
        return 0, 0, np.nan, np.nan, np.nan, np.nan
    cats = sorted(set(x) | set(y))
    idx = {c: i for i, c in enumerate(cats)}
    p = np.zeros((len(cats), len(cats)))
    for a, b in zip(x, y):
        p[idx[a], idx[b]] += 1
    agree = int(np.trace(p))
    p /= n
    po = np.trace(p)
    pr, pc = p.sum(axis=1), p.sum(axis=0)          # A 的边际、B 的边际
    pe = float(pr @ pc)
    if pe >= 1:                                    # 两人都只用了同一个类别：κ 无定义
        return n, agree, po, np.nan, np.nan, np.nan
    k = (po - pe) / (1 - pe)
    t1 = sum(p[i, i] * (1 - (pr[i] + pc[i]) * (1 - k)) ** 2 for i in range(len(cats)))
    t2 = (1 - k) ** 2 * sum(p[i, j] * (pc[i] + pr[j]) ** 2
                            for i in range(len(cats)) for j in range(len(cats)) if i != j)
    t3 = (k - pe * (1 - k)) ** 2
    se = math.sqrt(max(t1 + t2 - t3, 0) / n) / (1 - pe)
    return n, agree, po, k, max(-1.0, k - 1.96 * se), min(1.0, k + 1.96 * se)


def landis_koch(k):
    if pd.isna(k):
        return '不可估计'
    if k < 0:
        return '差'
    for cut, lab in [(0.20, '轻微'), (0.40, '一般'), (0.60, '中等'), (0.80, '高度')]:
        if k <= cut:
            return lab
    return '几乎完全'


def _r3(v):
    return round(v, 3) + 0.0 if pd.notna(v) else np.nan     # +0.0 去掉 -0.0


def agreement_table(a, b):
    rows = []
    for field, cond, desc in CAT_FIELDS + NUM_FIELDS:
        m = pd.Series(True, index=a.index) if cond is None else cond(a, b)
        m &= (a[field] != '') & (b[field] != '')
        xa, xb = a.loc[m, field], b.loc[m, field]
        note = ''
        if field in CAT_NAMES:
            n, agree, po, k, lo, hi = cohen_kappa(xa, xb)
            ua, ub = xa.nunique(), xb.nunique()
            if n == 0:
                note = '尚无可比病例'
            elif ua == 1 and ub == 1 and set(xa) == set(xb):
                note = '两人只用了同一类别，κ 无定义，以观察一致率为准'
            elif ua == 1 or ub == 1:
                lo = hi = np.nan               # 边际退化时大样本标准误不可用
                note = '一方只用了一个类别，κ 受边际分布限制（κ 悖论），以观察一致率为准'
        else:
            n, agree = int(m.sum()), int((xa == xb).sum())
            po, k, lo, hi = (agree / n if n else np.nan), np.nan, np.nan, np.nan
        rows.append({'字段': field, '计算范围': desc, '病例数': n, '一致数': agree,
                     '观察一致率': _r3(po), 'Cohen κ': _r3(k), '95% CI 下限': _r3(lo), '95% CI 上限': _r3(hi),
                     '一致程度（Landis-Koch）': ('仅报一致率' if field not in CAT_NAMES else
                                              '见备注' if note and pd.notna(k) else landis_koch(k)),
                     '备注': note})
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- 分歧与仲裁
def disagreements(a, b):
    rows = []
    for k in a.index:
        for f in ARB_FIELDS:
            va, vb = a.at[k, f], b.at[k, f]
            if va != vb:
                rows.append({KEY: k, '住院号': a.at[k, '住院号'], '入院日期': a.at[k, '入院日期'],
                             '年龄组': a.at[k, '年龄组'], '字段': f, '裁定者A': va, '裁定者B': vb,
                             '仲裁结果': '', '仲裁者': '', '仲裁理由': '', '回填提示': ''})
    return pd.DataFrame(rows, columns=[KEY, '住院号', '入院日期', '年龄组', '字段', '裁定者A', '裁定者B']
                        + ARB_COLS + ['回填提示'])


def carry_arbitration(dis):
    """把旧输出里已填写的仲裁结果按（就诊编号, 字段）回填，返回 (回填条数, 需重新仲裁条数)。

    若该条的 A/B 值已与上次运行时不同，旧仲裁不再沿用：仲裁列留空，旧仲裁内容写进“回填提示”，
    直到仲裁者重新填写为止（提示会随每次运行保留，不会自行消失）。"""
    if not os.path.exists(OUT) or dis.empty:
        return 0, 0
    old = pd.read_excel(OUT, sheet_name='分歧与仲裁', dtype=object, keep_default_na=False)
    if old.empty:
        return 0, 0
    for c in [KEY, '字段', '裁定者A', '裁定者B', '回填提示'] + ARB_COLS:
        old[c] = old[c].map(norm)
    old = old.drop_duplicates([KEY, '字段']).set_index([KEY, '字段'])
    n = stale = 0
    for i, r in dis.iterrows():
        key = (r[KEY], r['字段'])
        if key not in old.index:
            continue
        o = old.loc[key]
        if not any(o[c] for c in ARB_COLS):
            dis.at[i, '回填提示'] = o['回填提示']          # 尚待重新仲裁的提示继续保留
            continue
        if (o['裁定者A'], o['裁定者B']) != (r['裁定者A'], r['裁定者B']):
            dis.at[i, '回填提示'] = ('裁定已变更，请重新仲裁。原仲裁：%s（仲裁者 %s，理由：%s；当时 A=%s，B=%s）'
                                  % (o['仲裁结果'], o['仲裁者'], o['仲裁理由'], o['裁定者A'], o['裁定者B']))
            stale += 1
            continue
        for c in ARB_COLS:
            dis.at[i, c] = o[c]
        n += 1
    return n, stale


def final_dataset(a, b, dis):
    arb = {(r[KEY], r['字段']): r['仲裁结果'] for _, r in dis.iterrows() if r['仲裁结果']}
    out = a[[c for c in INFO_COLS if c in a.columns]].copy()
    done = (a['病因大类'] != '') & (b['病因大类'] != '')
    for f in ARB_FIELDS:
        vals = []
        for k in a.index:
            va, vb = a.at[k, f], b.at[k, f]
            vals.append(va if va == vb else arb.get((k, f), PENDING))
        out[f] = vals
    for f in TEXT_FIELDS:
        out[f] = [a.at[k, f] if a.at[k, f] == b.at[k, f] or not b.at[k, f] else
                  b.at[k, f] if not a.at[k, f] else 'A：%s｜B：%s' % (a.at[k, f], b.at[k, f])
                  for k in a.index]
    pend = (out[ARB_FIELDS] == PENDING).any(axis=1)
    out.insert(0, '终裁状态', np.where(~done, '裁定未完成', np.where(pend, '待仲裁', '已定稿')))
    return out


def confusion(a, b):
    m = (a['病因大类'] != '') & (b['病因大类'] != '')
    if not m.any():
        return pd.DataFrame()
    t = pd.crosstab(a.loc[m, '病因大类'], b.loc[m, '病因大类'],
                    rownames=['裁定者A'], colnames=['裁定者B'], margins=True, margins_name='合计')
    return t.reset_index()


# ---------------------------------------------------------------- 写盘
def write_out(info, checks, agr, conf, dis, fin):
    if os.path.exists(OUT):
        bak = OUT.replace('.xlsx', '.backup_%s.xlsx' % datetime.now().strftime('%Y%m%d_%H%M%S'))
        shutil.copy2(OUT, bak)
        print('已备份旧文件 ->', bak)
    bold = Font(bold=True)
    yellow = PatternFill('solid', fgColor='FFF2CC')
    red = PatternFill('solid', fgColor='FCE4D6')
    with pd.ExcelWriter(OUT, engine='openpyxl') as w:
        info.to_excel(w, sheet_name='说明', index=False)
        checks.to_excel(w, sheet_name='逻辑核查', index=False)
        agr.to_excel(w, sheet_name='一致性κ', index=False)
        conf.to_excel(w, sheet_name='病因大类交叉表', index=False)
        dis.to_excel(w, sheet_name='分歧与仲裁', index=False)
        fin.to_excel(w, sheet_name='终裁数据集', index=False)
        for name, ws in w.sheets.items():
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 16
            for c in ws[1]:
                c.font = bold
                c.alignment = Alignment(wrap_text=True, vertical='center')
            ws.freeze_panes = 'A2'
        ws = w.sheets['说明']
        ws.column_dimensions['A'].width = 22
        ws.column_dimensions['B'].width = 110
        for row in ws.iter_rows(min_row=2):
            row[1].alignment = Alignment(wrap_text=True, vertical='top')
        w.sheets['逻辑核查'].column_dimensions['F'].width = 70
        w.sheets['一致性κ'].column_dimensions['B'].width = 26

        # 仲裁列：黄色底，按字段加下拉校验
        ws = w.sheets['分歧与仲裁']
        cols = list(dis.columns)
        for c in ARB_COLS:
            j = cols.index(c) + 1
            ws.cell(row=1, column=j).fill = yellow
            for i in range(2, len(dis) + 2):
                ws.cell(row=i, column=j).fill = yellow
        jt = cols.index('回填提示') + 1
        ws.column_dimensions[ws.cell(row=1, column=jt).column_letter].width = 40
        for i, v in enumerate(dis['回填提示'], start=2):
            if v:
                ws.cell(row=i, column=jt).fill = red
        jr = cols.index('仲裁结果') + 1
        L = ws.cell(row=1, column=jr).column_letter
        for f, opts in adj.OPTIONS.items():
            rows = [i for i, v in enumerate(dis['字段'], start=2) if v == f]
            if not rows:
                continue
            dv = DataValidation(type='list', formula1='"%s"' % ','.join(opts), allow_blank=True)
            ws.add_data_validation(dv)
            for i in rows:
                dv.add('%s%d' % (L, i))

        ws = w.sheets['终裁数据集']
        for j, name in enumerate(fin.columns, start=1):
            if name in ARB_FIELDS:
                for i, v in enumerate(fin[name], start=2):
                    if v == PENDING:
                        ws.cell(row=i, column=j).fill = red
    print('已写出 ->', OUT)


def main():
    pa, pb = adj.rater_path('A'), adj.rater_path('B')
    a, b = load_rater(pa), load_rater(pb)
    ka, kb = set(a.index), set(b.index)
    assert ka == kb, '两份盲表病例不一致：仅A有 %s；仅B有 %s' % (sorted(ka - kb)[:5], sorted(kb - ka)[:5])
    b = b.loc[a.index]

    checks = pd.DataFrame(logic_check(a, 'A') + logic_check(b, 'B'),
                          columns=['裁定者', '级别', KEY, '住院号', '入院日期', '问题'])

    done_a, done_b = a['病因大类'] != '', b['病因大类'] != ''
    both = done_a & done_b
    agr = agreement_table(a[both], b[both])
    conf = confusion(a, b)

    dis = disagreements(a[both], b[both])
    n_carry, n_stale = carry_arbitration(dis)
    fin = final_dataset(a, b, dis)

    n_err = int((checks['级别'] == '错误').sum())
    n_arb = int((dis['仲裁结果'] != '').sum()) if not dis.empty else 0
    k_main = agr.loc[agr['字段'] == '病因大类'].iloc[0]
    info = pd.DataFrame([
        ('生成时间', datetime.now().strftime('%Y-%m-%d %H:%M')),
        ('裁定者A 盲表', '%s（SHA256 %s）' % (pa, sha256(pa))),
        ('裁定者B 盲表', '%s（SHA256 %s）' % (pb, sha256(pb))),
        ('进度', '共 %d 例；A 已完成 %d 例，B 已完成 %d 例，两人均完成 %d 例（κ 仅基于这些病例）'
                % (len(a), done_a.sum(), done_b.sum(), both.sum())),
        ('逻辑核查', '错误 %d 条、提示 %d 条。错误须退回本人在盲表中修改后重新运行本脚本；'
                    '带着错误算出的 κ 与分歧清单不可作为最终结果。'
                    % (n_err, int((checks['级别'] == '提示').sum()))),
        ('病因大类 κ', '%s（95%% CI %s–%s），n=%d，观察一致率 %s'
                       % (k_main['Cohen κ'], k_main['95% CI 下限'], k_main['95% CI 上限'],
                          k_main['病例数'], k_main['观察一致率'])),
        ('分歧与仲裁', '分歧 %d 条，已仲裁 %d 条（从旧文件回填 %d 条）；另有 %d 条因裁定已变更、旧仲裁作废需重新仲裁'
                     % (len(dis), n_arb, n_carry, n_stale)),
        ('终裁状态', '；'.join('%s %d 例' % (s, n) for s, n in fin['终裁状态'].value_counts().items())),
        ('方法说明', '病因判定由 2 名研究者依据统一的操作性定义（见盲表“分类标准”）独立完成，'
                    '裁定表不显示机器初筛提示。以未加权 Cohen κ 评价判定一致性，95% CI 按 Fleiss、Cohen 与 '
                    'Everitt（1969）大样本标准误计算；依赖前序判断的字段（排除理由、异物类型、穿孔部位等）'
                    '仅在两人对前序判断一致的病例中计算，避免“不适用对不适用”的一致抬高 κ。'
                    '分歧由第三名研究者参照原始病历仲裁。κ 基于仲裁前的独立判定。'),
        ('仲裁操作', '仲裁者只在“分歧与仲裁”表的黄色列填写，然后重新运行本脚本，“终裁数据集”即更新；'
                    '已填的仲裁结果每次运行都会自动回填，不会丢失。'),
    ], columns=['项目', '说明'])

    write_out(info, checks, agr, conf, dis, fin)
    print(info.iloc[3:8].to_string(index=False, header=False))
    print(agr[['字段', '病例数', '观察一致率', 'Cohen κ', '95% CI 下限', '95% CI 上限']].to_string(index=False))


if __name__ == '__main__':
    main()
