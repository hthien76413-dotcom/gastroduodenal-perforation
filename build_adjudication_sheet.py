# -*- coding: utf-8 -*-
"""
课题④ 儿童胃十二指肠穿孔 —— 生成双人裁定工作表
从 HIS 导出的 胃十二指肠穿孔.xlsx 汇总每例的判定依据文本，附空白裁定列。

输出三份文件：
  · 裁定工作表_v1.xlsx         总表，含机器初筛提示，仅供仲裁者参考；
  · 裁定工作表_v1_裁定者A.xlsx  裁定者 A 的盲表；
  · 裁定工作表_v1_裁定者B.xlsx  裁定者 B 的盲表。
盲表去掉「机器初筛提示」「磁性关键词命中」两列：两人若看到同一条提示，判定会被同一方向牵引，
κ 虚高，也就谈不上「独立」。两份盲表填完后用 merge_adjudication.py 计算 κ 并生成仲裁清单。

安全约定（吸取以往脚本覆盖已填 Excel 的教训）：
  1. 目标文件若已存在，先备份为 *.backup_时间戳.xlsx；
  2. 再按【科研就诊编号】把旧文件里已填写的裁定列逐条回填到新表，绝不清空人工劳动；
     A、B 两份盲表各自只从本人的旧盲表回填，互不串用；
  3. 所有写盘动作只在 __main__ 中触发，import 本模块不产生任何副作用。
"""
import os
import re
import shutil
from datetime import datetime

import pandas as pd
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill

SRC = r'D:\胃十二指肠穿孔\胃十二指肠穿孔.xlsx'
DST = r'D:\胃十二指肠穿孔\裁定工作表_v1.xlsx'
KEY = '科研就诊编号'

# 人工填写列 —— 回填与保护的对象
ADJ_COLS = [
    '病因大类', '异物类型', '磁性异物枚数', '异物具体名称',
    '穿孔部位', '合并其他消化道穿孔', '穿孔数目', '手术入路',
    'HP状态', '结局', '是否排除', '排除理由', '裁定者', '裁定日期', '备注',
]

# 下拉选项
OPTIONS = {
    '病因大类': ['1异物相关', '2新生儿自发性/胃壁肌层缺损', '3消化性溃疡',
                 '4创伤', '5医源性', '6其他/特发性', '9排除'],
    '异物类型': ['磁性', '非磁性', 'NA'],
    '穿孔部位': ['胃', '十二指肠', '胃+十二指肠'],
    '手术入路': ['开腹', '腹腔镜', '腔镜中转开腹', '内镜', '未手术'],
    'HP状态': ['阳性', '阴性', '未查'],
    '结局': ['治愈出院', '好转出院', '放弃治疗', '院内死亡'],
    '是否排除': ['否', '是'],
    '排除理由': ['E1部位非胃十二指肠', 'E2无活动性穿孔', 'E3外院术后转入',
                 'E4同次穿孔再住院', 'E5资料严重缺失'],
}

# 盲表中去掉的列：机器生成、直接指向病因的提示
HINT_COLS = ['机器初筛提示', '磁性关键词命中']
RATERS = ('A', 'B')

CRITERIA = [
    ('分类总则', '以手术经过 + 术中诊断为第一依据，现病史次之，出院诊断编码仅作参考。'
                 '两名裁定者独立填写、互不查看，完成后计算 Cohen κ，分歧由第三方仲裁。'),
    ('1 异物相关', '术中或内镜证实消化道内异物，且穿孔部位与异物滞留/嵌顿部位一致。'
                   '再按异物类型分磁性（磁珠、磁力珠、巴克球、磁铁）与非磁性（发卡、掏耳勺、电池等）；'
                   '磁性者须在“磁性异物枚数”登记术中取出的确切枚数，记不清写 NR。'),
    ('2 新生儿自发性/胃壁肌层缺损', '日龄≤28 天，术中见胃壁肌层缺损或胃壁自发破裂，'
                                    '无异物、无溃疡、无外伤与医源性操作史。'),
    ('3 消化性溃疡', '术中见慢性溃疡基底（胼胝样边缘、瘢痕化），或病理证实溃疡；'
                     '含十二指肠球部溃疡穿孔。HP 状态另列登记。'),
    ('4 创伤', '明确外力史（车祸、坠落、钝挫伤、虐待伤），术中见挫裂伤/系膜撕裂。'),
    ('5 医源性', '穿孔发生于诊疗操作后且部位吻合：内镜操作、置管、既往手术吻合口/修补处漏。'),
    ('6 其他/特发性', '不符合以上任一条，或术中未能确定病因。包括先天畸形继发'
                      '（十二指肠隔膜、环状胰腺）、NEC 累及、胃石、全身疾病背景下发生者，'
                      '须在备注写明具体情形，后续按频次决定是否单列亚类。'),
    ('9 排除', '符合下列任一条即排除，并在“排除理由”写明具体条目：\n'
               'E1 穿孔部位不在胃或十二指肠（如单纯空/回肠、结肠、食管穿孔）；\n'
               'E2 本次住院并无活动性穿孔（如仅为溃疡、术后复查、拔管或造影随访）；\n'
               'E3 穿孔已在外院手术处理后转入我院，本次住院处理的是并发症；\n'
               'E4 同一患儿因同一次穿孔的再次住院——保留首次收治的那次住院，'
               '后续住院排除（其信息可回补到首次记录的并发症/再手术字段）；\n'
               'E5 病案资料严重缺失：已回 HIS 调阅原始病历（入院记录、手术记录、出院记录、'
               '影像及内镜报告）后，仍无法获得判定病因所需的任何依据。'
               '仅因科研平台导出缺少出院小结或手术记录者，不得直接按 E5 排除。\n'
               '注意：保守治疗成功者与家属拒绝手术者**不属于排除**，应正常纳入并在'
               '“手术入路”填“未手术”。'),
    ('填写一致性', '病因大类填“9排除”时，是否排除必须填“是”并选择排除理由，反之亦然；'
                   '异物类型仅在病因大类为“1异物相关”时填磁性/非磁性，其余填 NA；'
                   '磁性异物须填枚数（正整数，记不清写 NR）；'
                   '纳入病例须填写穿孔部位、手术入路与结局。'
                   'merge_adjudication.py 会逐条核查，不一致之处列入“逻辑核查”表退回本人修改。'),
    ('数据完整性标记', '本列由脚本自动生成，标出无手术记录、无出院小结、非首次住院的病例。'
                       '带标记者请优先裁定：它们既是排除条目 E2-E4 的高发人群，'
                       '也可能是真实的非手术治疗病例，两者必须逐例读现病史区分，不得一律排除。'
                       '标“无出院小结”“无手术记录”者先回 HIS 调阅原始病历，'
                       '多为科研平台尚未同步，不等于病历缺失。'),
    ('穿孔部位', '按穿孔的解剖位置填写，仅限胃 / 十二指肠 / 两者兼有。'
                 '若同时存在空肠、回肠、结肠穿孔，另填“合并其他消化道穿孔”。'),
    ('结局', '“放弃治疗”指家属签字放弃后出院者，分析时与院内死亡合并为复合不良结局。'),
]


def _join(series, sep=' | '):
    vals = [str(v).strip() for v in series if pd.notna(v) and str(v).strip() not in ('', 'nan')]
    return sep.join(dict.fromkeys(vals))


def build_frame():
    """汇总每例的裁定依据，返回待填写的 DataFrame（纯计算，无副作用）。"""
    xl = pd.ExcelFile(SRC)
    head = xl.parse('病案首页基本信息').copy()
    head['入院日期'] = pd.to_datetime(head['入院日期'], errors='coerce')
    head['出院日期'] = pd.to_datetime(head['出院日期'], errors='coerce')

    def age_group(d):
        if pd.isna(d):
            return ''
        if d <= 28:
            return '1新生儿'
        if d < 365:
            return '2婴儿'
        if d < 365 * 3:
            return '3幼儿'
        if d < 365 * 6:
            return '4学龄前'
        return '5学龄期'

    df = pd.DataFrame({
        '科研患者编号': head['科研患者编号'],
        KEY: head[KEY],
        '住院号': head['住院号'],
        '第几次住院': head['第几次住院'],
        '性别': head['性别'],
        '年龄_岁': head['年龄（岁）'],
        '年龄_天': head['年龄（天）'],
        '年龄组': head['年龄（天）'].map(age_group),
        '入院日期': head['入院日期'].dt.strftime('%Y-%m-%d'),
        '入院年': head['入院日期'].dt.year,
        '住院天数': head['实际住院天数'],
        '入院科别': head['入院科别'],
        '门急诊诊断': head['门（急）诊诊断名称'],
    })

    # 出院诊断：主要诊断单列，全部诊断拼接
    dis = xl.parse('病案出院诊断编目后')
    is_main = dis['是否主要诊断'].astype(str).str.contains('是|1', na=False)
    df['主要出院诊断'] = df[KEY].map(dis[is_main].groupby(KEY)['诊断疾病名称'].apply(_join))
    df['全部出院诊断'] = df[KEY].map(dis.groupby(KEY)['诊断疾病名称'].apply(_join))

    # 手术记录：术中诊断 / 手术名称 / 手术经过（判定的第一依据）
    rec = xl.parse('住院病历手术记录')
    for col, name in [('术中诊断', '术中诊断'), ('手术名称', '手术名称'), ('手术经过', '手术经过')]:
        df[name] = df[KEY].map(rec.groupby(KEY)[col].apply(_join))
    rec_time = rec.copy()
    rec_time['t'] = pd.to_datetime(rec_time['手术日期及时间'], errors='coerce')
    df['手术日期'] = df[KEY].map(rec_time.groupby(KEY)['t'].min().dt.strftime('%Y-%m-%d %H:%M'))

    # 病案编码手术（辅助）
    op = xl.parse('病案手术操作编码后')
    df['编码手术(全部)'] = df[KEY].map(op.groupby(KEY)['手术操作名称'].apply(_join))

    # 主诉 / 现病史：儿科与新生儿科两张入院记录合并
    ped = xl.parse('儿科入院记录')
    neo = xl.parse('新生儿科入院记录')
    adm = pd.concat([ped[[KEY, '主诉', '现病史', '既往史']],
                     neo[[KEY, '主诉', '现病史', '既往史']]], ignore_index=True)
    for c in ['主诉', '现病史', '既往史']:
        df[c] = df[KEY].map(adm.groupby(KEY)[c].apply(_join))

    # 出院情况（判定结局）
    out = xl.parse('住院病历出院记录')
    df['出院情况'] = df[KEY].map(out.groupby(KEY)['出院情况'].apply(_join))
    df['诊疗经过'] = df[KEY].map(out.groupby(KEY)['诊疗经过'].apply(_join))

    # 数据完整性标记：标出需优先人工裁定的病例（可能触发排除条目 E2-E4）
    def flag(r):
        f = []
        if pd.isna(r['手术经过']) or str(r['手术经过']).strip() in ('', 'nan'):
            f.append('无手术记录')
        if pd.isna(r['出院情况']) or str(r['出院情况']).strip() in ('', 'nan'):
            f.append('无出院小结')
        if pd.notna(r['第几次住院']) and r['第几次住院'] > 1:
            f.append('第%s次住院' % int(r['第几次住院']))
        return '⚠' + '+'.join(f) if f else ''

    df['数据完整性标记'] = df.apply(flag, axis=1)

    # 机器初筛提示：只作参考，最终以人工裁定为准，不参与任何统计
    blob = (df['全部出院诊断'].fillna('') + ' ' + df['术中诊断'].fillna('') + ' '
            + df['手术名称'].fillna('') + ' ' + df['手术经过'].fillna('').str[:600])
    df['机器初筛提示'] = blob.map(_hint)
    df['磁性关键词命中'] = blob.map(
        lambda t: _join(pd.Series(re.findall(r'磁珠|磁力珠|磁铁|巴克球|磁性', str(t))), '/') or '')

    for c in ADJ_COLS:
        # 显式 object 列：pandas 3 会把全 '' 列推断为严格字符串类型，回填数字（枚数、穿孔数目）时报错
        df[c] = pd.Series('', index=df.index, dtype=object)

    # 把优先核查标记提到前列，避免被右侧长文本淹没
    cols = list(df.columns)
    cols.insert(cols.index('性别'), cols.pop(cols.index('数据完整性标记')))
    df = df[cols]

    return df.sort_values(['入院日期', KEY]).reset_index(drop=True)


def _hint(t):
    t = str(t)
    if re.search(r'异物|磁珠|磁力珠|磁铁|巴克球', t):
        return '?1异物相关'
    if re.search(r'胃壁肌层缺损', t):
        return '?2新生儿自发性'
    if re.search(r'溃疡|幽门螺', t):
        return '?3消化性溃疡'
    if re.search(r'损伤|外伤|车祸|坠落|挫伤', t):
        return '?4创伤'
    if re.search(r'术后|医源|瘘', t):
        return '?5医源性'
    return '?6其他/待读'


def rater_path(rater):
    """裁定者盲表路径，与总表同目录。"""
    return DST.replace('.xlsx', '_裁定者%s.xlsx' % rater)


def carry_over(df, path=None):
    """把旧文件里已填写的裁定列按就诊编号回填，返回回填条数。"""
    path = path or DST
    if not os.path.exists(path):
        return 0
    # keep_default_na=False：否则已填的“NA”（异物类型）会被当成缺失值、重建时丢失
    old = pd.read_excel(path, sheet_name='裁定表', keep_default_na=False)
    if KEY not in old.columns:
        return 0
    have = [c for c in ADJ_COLS if c in old.columns]
    old = old.set_index(KEY)[have]
    n = 0
    for i, k in df[KEY].items():
        if k in old.index:
            row = old.loc[k]
            if isinstance(row, pd.DataFrame):      # 同一 key 多行时取首行
                row = row.iloc[0]
            for c in have:
                v = row[c]
                if pd.notna(v) and str(v).strip() != '':
                    df.at[i, c] = v
                    n += 1
    return n


def write_workbook(df, path=None, rater=None):
    """rater 为 None 时写总表；为 'A'/'B' 时写该裁定者的盲表（调用方须已去掉 HINT_COLS）。"""
    path = path or DST
    assert rater is None or not set(HINT_COLS) & set(df.columns), '盲表不得含机器提示列'
    if os.path.exists(path):
        bak = path.replace('.xlsx', '.backup_%s.xlsx' % datetime.now().strftime('%Y%m%d_%H%M%S'))
        shutil.copy2(path, bak)
        print('已备份旧文件 ->', bak)

    if rater:
        role = ('裁定者 %s 的独立裁定盲表。请独立完成：勿与另一名裁定者讨论病例，勿查看对方表格或总表'
                '（总表含机器初筛提示）。填完后交课题负责人运行 merge_adjudication.py。' % rater)
    else:
        role = ('总表，含机器初筛提示，仅供仲裁者参考，请勿在本表中裁定。'
                '两名裁定者分别在 %s 与 %s 中独立填写。'
                % tuple(os.path.basename(rater_path(r)) for r in RATERS))
    crit = pd.DataFrame([('本表用途', role)] + CRITERIA, columns=['条目', '操作性定义'])
    with pd.ExcelWriter(path, engine='openpyxl') as w:
        crit.to_excel(w, sheet_name='分类标准', index=False)
        df.to_excel(w, sheet_name='裁定表', index=False)

        ws = w.sheets['裁定表']
        ncol = df.shape[1]
        fill = PatternFill('solid', fgColor='FFF2CC')       # 待填列底色
        for j, name in enumerate(df.columns, start=1):
            L = get_column_letter(j)
            if name in ADJ_COLS:
                ws.column_dimensions[L].width = 16
                ws.cell(row=1, column=j).fill = fill
            elif name in ('手术经过', '现病史', '诊疗经过', '出院情况', '全部出院诊断', '编码手术(全部)'):
                ws.column_dimensions[L].width = 60
            else:
                ws.column_dimensions[L].width = 14
            ws.cell(row=1, column=j).font = Font(bold=True)
            ws.cell(row=1, column=j).alignment = Alignment(wrap_text=True, vertical='center')
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = 'A1:%s1' % get_column_letter(ncol)

        # 带完整性标记的行整行提示色，提醒优先裁定
        jf = list(df.columns).index('数据完整性标记') + 1
        ws.column_dimensions[get_column_letter(jf)].width = 22
        warn = PatternFill('solid', fgColor='FCE4D6')
        for i, v in enumerate(df['数据完整性标记'], start=2):
            if str(v).strip():
                for j in range(1, ncol + 1):
                    ws.cell(row=i, column=j).fill = warn
                ws.cell(row=i, column=jf).font = Font(color='C00000', bold=True)

        # 下拉校验
        for name, opts in OPTIONS.items():
            if name not in df.columns:
                continue
            j = list(df.columns).index(name) + 1
            L = get_column_letter(j)
            dv = DataValidation(type='list', formula1='"%s"' % ','.join(opts), allow_blank=True)
            ws.add_data_validation(dv)
            dv.add('%s2:%s%d' % (L, L, len(df) + 1))

        cw = w.sheets['分类标准']
        cw.column_dimensions['A'].width = 30
        cw.column_dimensions['B'].width = 110
        for r in range(1, len(crit) + 2):
            cw.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical='top')
        cw.cell(row=2, column=2).font = Font(bold=True, color='C00000')
    print('已写出 ->', path)


if __name__ == '__main__':
    base = build_frame()

    d = base.copy()
    n = carry_over(d)
    print('汇总 %d 例；总表从旧文件回填已填写单元格 %d 个' % (len(d), n))
    write_workbook(d)

    for r in RATERS:
        b = base.drop(columns=HINT_COLS)          # 每份盲表都从空白起步，只回填本人旧表
        n = carry_over(b, rater_path(r))
        print('裁定者 %s 盲表：从本人旧表回填已填写单元格 %d 个' % (r, n))
        write_workbook(b, rater_path(r), rater=r)

    print('\n机器初筛提示分布（仅参考，不入统计）:')
    print(base['机器初筛提示'].value_counts().to_string())
