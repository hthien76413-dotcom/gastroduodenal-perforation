# -*- coding: utf-8 -*-
"""生成《病案数据调取_填报模板》——病案室直接在格子里填数，避免口径偏差。

安全约定（同 build_adjudication_sheet.py）：目标文件若已存在先备份，写盘只在 __main__。
本文件是发给病案室的空白模板；若病案室已回填，请勿直接重跑覆盖。
"""
import os
import shutil
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

DST = r'D:\胃十二指肠穿孔\①儿童胃十二指肠穿孔病因谱十年变迁\病案数据调取_填报模板_v1.xlsx'

BATCH = '2026R090-E01'
Y0, M0, Y1, M1 = 2016, 1, 2026, 5

# 必需病区（本课题病例的主要来源，覆盖 95/107 例）
WARD_CORE = ['新生儿外科病区', '普外一病区', '普外二病区',
             '胃肠外科病区', '肝胆外科/肿瘤外科病区']
# 一并提供，用于敏感性分析（另覆盖 10 例）
WARD_EXT = ['新生儿内科一病区', '新生儿内科二病区', '新生儿内科病区(西院)', '重症医学科病区']
WARDS = WARD_CORE + WARD_EXT

HDR_FILL = PatternFill('solid', fgColor='D9E1F2')
CORE_FILL = PatternFill('solid', fgColor='FFF2CC')     # 必需，待填
EXT_FILL = PatternFill('solid', fgColor='E2EFDA')      # 敏感性，待填
NOTE_FILL = PatternFill('solid', fgColor='FCE4D6')
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def months():
    out, y, m = [], Y0, M0
    while (y, m) <= (Y1, M1):
        out.append((y, m))
        m += 1
        if m == 13:
            y, m = y + 1, 1
    return out


def style_header(ws, row, ncol):
    for j in range(1, ncol + 1):
        c = ws.cell(row=row, column=j)
        c.font = Font(bold=True)
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BOX


def sheet_note(ws, lines):
    """在表顶部写若干行说明，返回表头应放置的行号。"""
    r = 1
    for text, bold in lines:
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(bold=bold, size=12 if bold else 10)
        c.alignment = Alignment(vertical='center')
        r += 1
    return r + 1


def build_matrix_sheet(wb, title, notes, wards, extra_cols=None):
    """年月 × 病区 的矩阵式填报表。"""
    ws = wb.create_sheet(title)
    hr = sheet_note(ws, notes)
    cols = ['年', '月'] + list(wards) + list(extra_cols or [])
    for j, name in enumerate(cols, start=1):
        ws.cell(row=hr, column=j, value=name)
    style_header(ws, hr, len(cols))
    for i, (y, m) in enumerate(months()):
        r = hr + 1 + i
        ws.cell(row=r, column=1, value=y).border = BOX
        ws.cell(row=r, column=2, value=m).border = BOX
        for j in range(3, len(cols) + 1):
            c = ws.cell(row=r, column=j)
            c.border = BOX
            c.fill = EXT_FILL if cols[j - 1] in WARD_EXT else CORE_FILL
    ws.column_dimensions['A'].width = 7
    ws.column_dimensions['B'].width = 5
    for j in range(3, len(cols) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 17
    ws.freeze_panes = ws.cell(row=hr + 1, column=3)
    return ws, hr


def build_info_sheet(wb):
    ws = wb.create_sheet('填表说明')
    rows = [
        ('武汉儿童医院病案室 数据调取填报模板', None),
        ('', None),
        ('课题名称', '儿童胃十二指肠穿孔病因谱的十年变迁及临床特征分析'),
        ('伦理批件号', BATCH),
        ('申请人 / 科室', '舒俊 / 胃肠外科'),
        ('联系电话', '13545265102'),
        ('', None),
        ('一、统一口径（各表共用，务必一致）', None),
        ('时间范围', '2016 年 1 月至 2026 年 5 月，共 125 个月；模板中已预置全部年月行'),
        ('归属月份', '按【入院日期】归属，不按出院日期；跨月住院计入入院当月'),
        ('计数单位', '【住院人次】——每次住院计 1 次；同一患儿多次住院分别计入，不按患者去重'),
        ('科别口径', '按【出院科别】统计'),
        ('空值处理', '某病区某月无住院请填 0；不要留空，也不要删除整行'),
        ('数值格式', '纯数字，不带千分位分隔符与单位'),
        ('', None),
        ('二、各表填报要求', None),
        ('第1项（必需）', '外科病区逐月住院人次——本课题主分母。黄色列为必需，绿色列用于敏感性分析'),
        ('第2项（若支持）', '同第1项，但仅统计【急诊入院】者；若系统无此字段，请在该表注明后跳过'),
        ('第3项（必需）', '消化道异物（ICD-10 T18 及其全部亚目）逐月住院人次'),
        ('第4项（必需）', '病区沿革、统计字段与数据完整性确认——五个问题，请直接在表内作答'),
        ('第5项（可选）', '全院逐月住院总人次，仅用于敏感性分析'),
        ('', None),
        ('三、重要提示', None),
        ('病区更名', '据了解，2024 年 8 月前后普外一病区更名为胃肠外科病区、普外二病区更名为'
                     '肝胆外科/肿瘤外科病区，在此之前该两病区并未设立。'),
        ('', '更名前后的四个名称在模板中均已单独列出，请按各年月实际情况分别填写，不要预先合并。'
             '本课题将在分析时自行归并，因此某名称在某些年份无数据时请填 0——这一信息对我们同样重要。'),
        ('名称变体', '若贵室系统中另有「胃肠外科」「肝胆外科/肿瘤外科」等不带「病区」二字的写法，'
                     '请一并计入对应列。'),
        ('2019-07～2020-03', '本课题的科研平台导出在这 9 个月为空。请按贵室病案首页系统的实际数字照常填报，'
                             '不要填 0；若贵室系统同样缺失，请在第 4 项注明。'),
        ('保密说明', '本次调取均为汇总计数，不含任何患者可识别信息，仅用于上述已获伦理批准课题的'
                     '分母估计与发生率分析。'),
    ]
    r = 1
    for a, b in rows:
        if b is None:
            c = ws.cell(row=r, column=1, value=a)
            c.font = Font(bold=bool(a), size=13 if a.startswith(('武汉', '一、', '二、', '三、')) else 11)
        else:
            ws.cell(row=r, column=1, value=a).font = Font(bold=True)
            c = ws.cell(row=r, column=2, value=b)
            c.alignment = Alignment(wrap_text=True, vertical='center')
            ws.row_dimensions[r].height = 30 if len(b) > 55 else 16
        r += 1
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 92
    return ws


def build_qa_sheet(wb):
    ws = wb.create_sheet('第4项_病区沿革确认')
    known = ('已知情况（由课题组提供，请贵室核实）：2024 年 8 月前后，普外一病区更名为胃肠外科病区，'
             '普外二病区更名为肝胆外科/肿瘤外科病区；在此之前该两病区并未设立。')
    qs = [
        '1. 上述更名的确切生效日期（至少精确到月）。',
        '2. 贵室统计住院量所依据的字段是【病区】还是【专业组】？'
        '本课题病例资料中，2020 至 2023 年已有部分病例的出院科别显示为「胃肠外科病区」或'
        '「肝胆外科/肿瘤外科病区」，而该两病区当时尚未设立，推测这些记录反映的是普外科内部的'
        '专业组归属而非实际病区。请说明两者在历史数据中是否存在类似不一致。',
        '3. 历史住院量是否已按现行病区名称追溯归并？若已追溯，请注明追溯的起始时间；'
        '若未追溯，则 2024 年 8 月以前的住院量应仍挂在普外一病区、普外二病区名下。',
        '4. 2019 年 7 月至 2020 年 3 月的住院数据是否完整？本课题从院内科研数据平台导出的两份数据'
        '（胃十二指肠穿孔、消化道异物）在这 9 个月内住院记录均为零，而前后月份均正常'
        '（消化道异物住院每月 5 至 15 次），推测为平台数据缺失。请说明：（1）贵室病案首页系统在此期间的'
        '住院数据是否完整；（2）若完整，第 1、2、3、5 项请照常填报这 9 个月的实际数字，不要因科研平台'
        '缺失而填 0；（3）如知晓科研平台缺失的原因（如系统切换、数据迁移）及能否补导，请一并说明。',
        '5. 其他需要说明的情况（如病区搬迁、床位数重大调整、统计规则变更等）。',
    ]
    ws.cell(row=1, column=1, value='第 4 项（必需）病区沿革与统计字段确认').font = Font(bold=True, size=13)
    c = ws.cell(row=3, column=1, value=known)
    c.alignment = Alignment(wrap_text=True, vertical='center')
    c.fill = NOTE_FILL
    c.border = BOX
    ws.row_dimensions[3].height = 46

    ws.cell(row=5, column=1, value='问题')
    ws.cell(row=5, column=2, value='贵室答复')
    style_header(ws, 5, 2)

    r = 6
    for q in qs:
        cq = ws.cell(row=r, column=1, value=q)
        cq.alignment = Alignment(wrap_text=True, vertical='top')
        cq.border = BOX
        ca = ws.cell(row=r, column=2)
        ca.fill = CORE_FILL
        ca.border = BOX
        ca.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 78 if len(q) > 60 else 40
        r += 1

    r += 1
    for lab in ['经办人签字', '完成日期']:
        ws.cell(row=r, column=1, value=lab).font = Font(bold=True)
        ws.cell(row=r, column=2).fill = CORE_FILL
        ws.cell(row=r, column=2).border = BOX
        r += 1
    ws.column_dimensions['A'].width = 74
    ws.column_dimensions['B'].width = 48
    return ws


def main():
    if os.path.exists(DST):
        bak = DST.replace('.xlsx', '.backup_%s.xlsx' % datetime.now().strftime('%Y%m%d_%H%M%S'))
        shutil.copy2(DST, bak)
        print('已备份旧文件 ->', bak)

    wb = Workbook()
    wb.remove(wb.active)

    build_info_sheet(wb)

    build_matrix_sheet(
        wb, '第1项_外科病区住院人次',
        [('第 1 项（必需）逐月 × 逐病区 住院人次', True),
         ('黄色列＝必需（本课题病例主要来源，覆盖 95/107 例）；绿色列＝一并提供，用于敏感性分析。', False),
         ('按入院日期归月；单位为住院人次；某病区某月无住院请填 0，勿留空或删行。', False)],
        WARDS)

    build_matrix_sheet(
        wb, '第2项_急诊入院住院人次',
        [('第 2 项（若系统支持）逐月 × 逐病区 住院人次——仅限【急诊入院】', True),
         ('筛选条件：入院途径（或入院方式）为「急诊」。若无此字段或不可靠，'
          '请在此注明后跳过本表：________________', False),
         ('其余口径同第 1 项。', False)],
        WARDS)

    ws, hr = build_matrix_sheet(
        wb, '第3项_消化道异物住院人次',
        [('第 3 项（必需）消化道异物 逐月住院人次', True),
         ('筛选：出院诊断（主要诊断或其他诊断中任一出现即计入）编码属于 ICD-10 T18 及其全部亚目。', False),
         ('本课题病例中实际出现过的编码，供核对贵室编码习惯：'
          'T18.100、T18.200、T18.300、T18.300x003、T18.301、T18.400、T18.900、T18.901。', False)],
        [], extra_cols=['全院 T18 住院人次', '其中主要诊断为 T18 的人次', '其中收治于第1项各病区的人次'])
    for j in (3, 4, 5):
        ws.column_dimensions[get_column_letter(j)].width = 27

    build_qa_sheet(wb)

    ws, hr = build_matrix_sheet(
        wb, '第5项_全院住院总人次',
        [('第 5 项（可选）全院逐月住院总人次', True),
         ('仅用作敏感性分析。该指标在新冠流行期及 2023 年支原体肺炎流行期受呼吸系统疾病'
          '住院量剧烈波动影响，本课题不将其作为主要分母。若调取不便可跳过。', False),
         ('', False)],
        [], extra_cols=['全院住院人次'])
    ws.column_dimensions['C'].width = 20

    wb.save(DST)
    print('已写出 ->', DST)
    print('工作表:', wb.sheetnames)
    print('每张矩阵表 %d 个年月行' % len(months()))


if __name__ == '__main__':
    main()
