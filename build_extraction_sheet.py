# -*- coding: utf-8 -*-
"""
儿童胃十二指肠穿孔 —— 生成次要终点提取表（1 人提取、第 2 人逐例核对）

对应方案 V2.0 次要终点与《统计分析计划_SAP_v1》3.6：临床表现、住院期间并发症（Clavien–Dindo 分级）、
非计划再次手术/再干预、恢复肠内进食时间。这些变量不在裁定表内，也不需要双人独立判定。

表内三类列：
  · 自动列（灰色表头）：由脚本从科研平台导出计算，如术后住院天数、首次手术后的全麻操作、切口愈合等级、
    入院后新发诊断、治疗前影像的游离气体句段、主诉病程预估值。只作提取时的线索，不直接进入分析；
  · 人工列（黄色表头）：提取者按“变量定义”填写，核对者逐例复核；
  · 原文列：主诉、现病史、体格检查、手术记录、诊疗经过、出院情况、出院医嘱等，供阅读。
表中不显示病因相关的机器提示，也不显示病因裁定结果，可与双人裁定同时进行而不破坏其盲法。

注意：科研平台未导出病程记录。进食时间在出院记录中多数未写明，须回 HIS 查阅病程记录，
查不到填 NR，不得推测。

安全约定（同 build_adjudication_sheet.py）：目标文件若已存在先备份；按科研就诊编号回填已填写的人工列，
绝不清空人工劳动；“逻辑核查”与“进度”两张表每次运行按当前已填内容重算；写盘只在 __main__。
"""
import os
import re
import shutil
from datetime import datetime

import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from merge_adjudication import norm

SRC = r'D:\胃十二指肠穿孔\胃十二指肠穿孔.xlsx'
MERGED = r'D:\胃十二指肠穿孔\裁定合并_一致性与仲裁_v1.xlsx'
DST = r'D:\胃十二指肠穿孔\次要终点提取表_v1.xlsx'
KEY = '科研就诊编号'

YN3 = ['是', '否', '不详']
COMP = ['否', '是', '不详', '不适用']
CLAVIEN = ['0', 'I', 'II', 'IIIa', 'IIIb', 'IVa', 'IVb', 'V']
SYMPTOMS = ['腹痛', '呕吐', '发热', '腹胀', '腹膜刺激征', '入院时休克或脓毒症']
COMPLICATIONS = ['切口感染或裂开', '腹腔脓肿或腹腔感染', '修补口漏或再穿孔', '术后肠梗阻',
                 '新发脓毒症或感染性休克', '肺部感染', '消化道出血']
FEED = ['首次肠内摄入_第几天', '全量肠内营养_第几天']
META = ['信息来源', '提取者', '提取日期', '核对者', '核对日期', '核对修改说明', '备注']
MANUAL_COLS = (['主诉病程_小时'] + SYMPTOMS + ['影像游离气体_治疗前'] + COMPLICATIONS
               + ['其他并发症', '最高Clavien-Dindo分级', '非计划再次手术', '再次手术原因', '非手术再干预']
               + FEED + META)
OPTIONS = dict({s: YN3 for s in SYMPTOMS}, **{c: COMP for c in COMPLICATIONS})
OPTIONS.update({'影像游离气体_治疗前': ['有', '无', '可疑', '未查'],
                '最高Clavien-Dindo分级': CLAVIEN,
                '非计划再次手术': ['否', '是', '不适用'],
                '非手术再干预': ['否', '是'],
                '信息来源': ['科研导出', '科研导出+HIS补查']})
# 必须填写的人工列（其余为条件必填或选填）
REQUIRED = ['主诉病程_小时'] + SYMPTOMS + ['影像游离气体_治疗前'] + COMPLICATIONS + \
           ['最高Clavien-Dindo分级', '非计划再次手术', '非手术再干预'] + FEED + ['信息来源']

DEFS = [
    ('本表用途', '次要终点的单人提取＋第二人逐例核对。人工列（黄色）按下列定义填写；灰色自动列只作线索，'
                 '须对照原文确认，不得照抄。终裁判为排除的病例整行灰显，无需提取。'),
    ('数据来源', '入院记录、手术记录、出院记录（诊疗经过、出院情况、出院医嘱）与影像报告。科研平台未导出病程记录：'
                 '原文未记载的项目回 HIS 查阅病程记录，并将“信息来源”填为“科研导出+HIS补查”；仍查不到的填“不详”或 NR，不得推测。'),
    ('计时起点', '手术病例以首次手术日为第 0 天（术后第 N 天）；未手术病例以入院日为第 0 天。自动列“进食计时起点”已给出日期。'),
    ('主诉病程_小时', '本次急性起病（出现与穿孔相关的症状，或吞食异物）至入院的时长，换算为小时（1 天＝24 小时，'
                      '“半天”＝12 小时）。主诉含慢性病史（如“间断腹痛 1 年……加重 2 天”）时取急性加重的时长；'
                      '异物病例按吞食时长计并在备注注明。自动列“预估_主诉病程_小时”仅供参考。无法确定填 NR。'),
    ('症状（腹痛、呕吐、发热、腹胀）', '本次病程中入院前或入院时出现即为“是”。发热指记录体温≥38.0℃或明确记载“发热”。'
                                    '新生儿及不能表达的患儿，腹痛填“不详”。'),
    ('腹膜刺激征', '入院查体记录压痛伴反跳痛，或腹肌紧张、板状腹。新生儿以记录的腹壁紧张、腹壁红肿为准。'),
    ('入院时休克或脓毒症', '入院记录、入院诊断中有休克、脓毒症（脓毒血症）、感染性休克，或入院 24 小时内使用血管活性药物。'),
    ('影像游离气体_治疗前', '首次手术前（未手术者入院 48 小时内）的 X 线或 CT 报告提示腹腔或膈下游离气体为“有”；'
                            '“不除外”“可疑”为“可疑”；相关报告均为未见游离气体为“无”；无 X 线或 CT 检查为“未查”。超声不计。'),
    ('并发症（通则）', '首次手术后（未手术者开始治疗后）本次住院期间新发的事件；入院时已存在的状况（如入院时已有的脓毒症、'
                       '腹腔感染）不计。每项填 是/否/不详；与治疗方式无关的项目（如未手术者的切口感染）填“不适用”。'),
    ('切口感染或裂开', '切口红肿渗液、化脓、脂肪液化或裂开；自动列中切口愈合等级为乙、丙，或出院诊断含 T81 者须重点核对。'),
    ('腹腔脓肿或腹腔感染', '影像证实的腹腔积脓或脓肿，或临床诊断的新发腹腔感染。'),
    ('修补口漏或再穿孔', '影像、引流物或再次手术证实的修补口漏、吻合口漏或新发穿孔。'),
    ('术后肠梗阻', '临床诊断的肠梗阻，且需禁食、胃肠减压或手术处理。'),
    ('新发脓毒症或感染性休克', '入院时不存在、治疗后新诊断的脓毒症或感染性休克。'),
    ('肺部感染', '临床诊断并给予抗感染治疗的肺炎。'),
    ('消化道出血', '呕血、便血或胃管引出血性液，需要处理者。'),
    ('其他并发症', '上列以外的并发症，以文字写明；无则留空。'),
    ('最高Clavien-Dindo分级', '按 Dindo 2004 标准取本次住院最高级别：0 无并发症；I 无需药物、手术或介入处理'
                              '（允许止吐、退热、镇痛、利尿、补电解质及床旁开放切口）；II 需药物治疗（含输血、全胃肠外营养、'
                              '因并发症调整抗生素）；IIIa 非全麻下的介入、内镜或穿刺；IIIb 全麻下的手术或干预；'
                              'IVa 单器官功能障碍需 ICU 处理；IVb 多器官功能障碍；V 死亡。针对穿孔本身的常规抗感染不算并发症处理。'),
    ('非计划再次手术', '本次住院期间因并发症或治疗失败进行的、事先未计划的全麻手术；事先计划的分期手术不计。'
                       '未手术病例填“不适用”。填“是”时写明再次手术原因。自动列“提示_再次全麻手术”非空者须重点核对。'),
    ('非手术再干预', '本次住院期间因并发症进行的穿刺引流、内镜或介入治疗。'),
    ('首次肠内摄入_第几天', '首次经口或经管饲摄入（含饮水、糖水、奶）的日期距计时起点的天数；原文与 HIS 均未记载填 NR。'),
    ('全量肠内营养_第几天', '停用静脉营养、经口或管饲达全量（新生儿按医嘱全量奶，年长儿恢复半流质或普食）的天数；未记载填 NR。'),
    ('核对', '核对者须与提取者不同。核对者逐例复核全部人工列，直接改正错误，并在“核对修改说明”写明改动的字段与原值；'
             '无改动写“无”。填写核对者与核对日期即视为该例核对完成。'),
]

TEXT_COLS = ['主诉', '现病史', '体格检查', '手术记录', '诊疗经过', '出院情况', '出院医嘱', '全部出院诊断']
HINT_WIDE = ['编码手术清单', '提示_影像游离气体句段', '提示_入院后新发诊断']

_CN = {'半': 0.5, '一': 1, '两': 2, '二': 2, '三': 3, '四': 4, '五': 5, '六': 6, '七': 7, '八': 8, '九': 9, '十': 10}
_UNIT_H = {'分钟': 1 / 60, '小时': 1, 'h': 1, '天': 24, '日': 24, '周': 168, '月': 720, '年': 8760}
_DUR = re.compile(r'(\d+(?:\.\d+)?|[半一两二三四五六七八九十]+)\s*(?:余|多)?\s*(?:个)?\s*(分钟|小时|h|天|日|周|月|年)')
_FA_SENT = re.compile(r'[^。；\n]*(?:游离气体|气腹|膈下游离)[^。；\n]*')
_FA_CLAUSE = re.compile(r'[^。；，,\n]*(?:游离气体|气腹|膈下游离)[^。；，,\n]*')
_FA_NEG = re.compile(r'未见|未显示|无明显|无游离|没有|消失|吸收')
_FA_EQV = re.compile(r'不除外|可疑|可能|待排')
# 手术记录的“手术日期及时间”只有日期；诊疗经过中常写有精确时刻，如“于2017-6-25 02:20:26全麻下行……”
_OP_TIME = re.compile(r'(\d{4}-\d{1,2}-\d{1,2})\s*(\d{1,2}:\d{2}(?::\d{2})?)[^，。；]{0,15}?(?:全麻|全身麻醉|麻醉)')
_POSTOP = re.compile(r'术后|引流管')


def _cn_num(s):
    if s[0].isdigit():
        return float(s)
    if s == '半':
        return 0.5
    if '十' in s:
        tens = _CN.get(s[0], 1) if s[0] != '十' else 1
        ones = _CN.get(s[-1], 0) if s[-1] != '十' else 0
        return tens * 10 + ones
    return _CN.get(s[0])


def parse_duration_h(text):
    """主诉病程预估（小时）：含“加重”时取加重后的时长，否则取最长一段；解析不到返回空。"""
    text = str(text or '')
    worse = re.search('加重[^，,。；]*', text)
    if worse and _DUR.search(worse.group(0)):
        text = worse.group(0)
    vals = []
    for num, unit in _DUR.findall(text):
        n = _cn_num(num)
        if n is not None:
            vals.append(n * _UNIT_H[unit])
    return round(max(vals), 1) if vals else None


def free_air(snips):
    """治疗前 X 线/CT 句段 → 有 / 可疑 / 无 / 未提及。否定词只在含关键词的分句内判断，
    以免“见游离气体，未见液平面”被后一分句的“未见”误判为阴性。"""
    clauses = [c for s in (snips or []) for c in _FA_CLAUSE.findall(s)]
    if not clauses:
        return '未提及'
    if any(not _FA_NEG.search(c) and not _FA_EQV.search(c) for c in clauses):
        return '有'
    if any(_FA_EQV.search(c) and not _FA_NEG.search(c) for c in clauses):
        return '可疑'
    return '无'


def _join(series, sep=' | '):
    vals = [str(v).strip() for v in series if pd.notna(v) and str(v).strip() not in ('', 'nan')]
    return sep.join(dict.fromkeys(vals))


def auto_frame(src=None):
    """每次住院的自动计算列（纯计算，无副作用）。analyze_etiology.py 也调用本函数。"""
    x = pd.ExcelFile(src or SRC)
    head = x.parse('病案首页基本信息')
    head[KEY] = head[KEY].map(norm)
    head['入院'] = pd.to_datetime(head['入院日期'], errors='coerce')
    head['出院'] = pd.to_datetime(head['出院日期'], errors='coerce')
    a = head[[KEY, '科研患者编号', '入院', '出院']].copy().set_index(KEY)
    a['住院天数'] = head.set_index(KEY)['实际住院天数']

    rec = x.parse('住院病历手术记录')
    rec[KEY] = rec[KEY].map(norm)
    rec['t'] = pd.to_datetime(rec['手术日期及时间'], errors='coerce')
    op = x.parse('病案手术操作编码后')
    op[KEY] = op[KEY].map(norm)
    op['t'] = pd.to_datetime(op['手术操作日期'], errors='coerce')
    ga = op[op['麻醉方式'].astype(str).str.contains('全') & op['t'].notna()]   # 全麻下的编码操作
    # 首次手术日期：手术记录优先，无手术记录时取编码手术日期（两者都只精确到日）
    a['首次手术时间'] = (rec.groupby(KEY)['t'].min().combine_first(ga.groupby(KEY)['t'].min())
                     .dt.normalize().reindex(a.index))
    out = x.parse('住院病历出院记录')
    out[KEY] = out[KEY].map(norm)
    exact = {}
    for k, t in zip(out[KEY], out['诊疗经过'].fillna('')):
        ts = pd.to_datetime(['%s %s' % m for m in _OP_TIME.findall(t)], errors='coerce').dropna()
        if len(ts) and k in a.index and pd.notna(a.at[k, '首次手术时间']) \
                and abs(ts.min().normalize() - a.at[k, '首次手术时间']) <= pd.Timedelta(days=1):
            exact[k] = ts.min()
    a['首次手术时刻'] = pd.Series(exact, dtype='datetime64[ns]').reindex(a.index)
    a['有手术记录'] = a['首次手术时间'].notna().map({True: '是', False: '否'})
    a['术后住院天数'] = (a['出院'].dt.normalize() - a['首次手术时间'].dt.normalize()).dt.days
    op['条目'] = (op['t'].dt.strftime('%Y-%m-%d').fillna('日期不详') + ' ' + op['手术操作名称'].fillna('').astype(str)
                + op['切口愈合等级'].map(lambda v: '［切口%s］' % v if pd.notna(v) and v not in ('其他', '无切口') else ''))
    a['编码手术清单'] = op.sort_values('t').groupby(KEY)['条目'].apply(_join).reindex(a.index).fillna('')
    # 首个全麻日期之后的全麻操作：非计划再次手术的线索（置管、呼吸机等非全麻操作不计）
    ga = ga.assign(d=ga['t'].dt.normalize())
    later = ga[ga['d'] > ga[KEY].map(ga.groupby(KEY)['d'].min())]
    a['提示_再次全麻手术'] = later.assign(
        条目=later['t'].dt.strftime('%Y-%m-%d') + ' ' + later['手术操作名称'].fillna('').astype(str)
    ).groupby(KEY)['条目'].apply(_join).reindex(a.index).fillna('')
    a['进食计时起点'] = [('首次手术日 %s' % f.strftime('%Y-%m-%d')) if pd.notna(f) else ('入院日 %s' % d.strftime('%Y-%m-%d'))
                   for f, d in zip(a['首次手术时间'], a['入院'])]

    dx = x.parse('病案出院诊断编目后')
    dx[KEY] = dx[KEY].map(norm)
    a['全部出院诊断'] = dx.groupby(KEY)['诊断疾病名称'].apply(_join).reindex(a.index).fillna('')
    new = dx[dx['入院病情'].astype(str).str.strip() == '无']
    a['提示_入院后新发诊断'] = new.groupby(KEY)['诊断疾病名称'].apply(_join).reindex(a.index).fillna('')
    t81 = dx[dx['诊断疾病编码'].astype(str).str.upper().str.startswith('T81')]
    wound = op[op['切口愈合等级'].isin(['乙', '丙'])]
    a['提示_切口'] = [
        '；'.join(filter(None, ['切口愈合等级乙/丙' if k in set(wound[KEY]) else '',
                                'T81：' + _join(t81.loc[t81[KEY] == k, '诊断疾病名称'], '、') if k in set(t81[KEY]) else '']))
        for k in a.index]

    # 治疗前（首次手术前；未手术者入院 48 小时内）X 线与 CT 报告中的游离气体句段
    img = []
    for sheet, typ in (('X线报告', 'X线'), ('CT报告', 'CT')):
        r = x.parse(sheet)
        r[KEY] = r[KEY].map(norm)
        r['t'] = pd.to_datetime(r['检查时间'], errors='coerce')
        r['文本'] = r['检查所见'].fillna('').astype(str) + '。' + r['检查结论'].fillna('').astype(str)
        r['类型'] = typ
        img.append(r[[KEY, 't', '文本', '类型']])
    img = pd.concat(img)
    # 治疗前窗口：有精确手术时刻者截至该时刻；只有手术日期者截至当日末并剔除写有“术后”“引流管”的报告；
    # 未手术者截至入院后 48 小时
    img['入院'] = img[KEY].map(a['入院'])
    exact_t = img[KEY].map(a['首次手术时刻'])
    day_end = img[KEY].map(a['首次手术时间']) + pd.Timedelta(hours=23, minutes=59)
    img['止'] = exact_t.fillna(day_end).fillna(img['入院'] + pd.Timedelta(hours=48))
    postop = exact_t.isna() & day_end.notna() & img['文本'].str.contains(_POSTOP)
    img = img[(img['t'] >= img['入院'] - pd.Timedelta(days=1)) & (img['t'] <= img['止']) & ~postop]
    snip, n_img = {}, img.groupby(KEY).size()
    for _, r in img.sort_values('t').iterrows():
        for s in _FA_SENT.findall(r['文本']):
            snip.setdefault(r[KEY], []).append('%s %s：%s' % (r['t'].strftime('%m-%d %H:%M'), r['类型'], s.strip()))
    a['提示_影像游离气体'] = [free_air(snip.get(k)) if n_img.get(k, 0) else '无X线/CT' for k in a.index]
    a['提示_影像游离气体句段'] = [' | '.join(snip.get(k, [])) for k in a.index]

    # 主诉病程预估
    ped, neo = x.parse('儿科入院记录'), x.parse('新生儿科入院记录')
    adm = pd.concat([ped[[KEY, '主诉', '现病史', '体格检查']], neo[[KEY, '主诉', '现病史', '体格检查']]])
    adm[KEY] = adm[KEY].map(norm)
    for c in ('主诉', '现病史', '体格检查'):
        a[c] = adm.groupby(KEY)[c].apply(_join).reindex(a.index).fillna('')
    a['预估_主诉病程_小时'] = a['主诉'].map(parse_duration_h)

    # 检索集内同一患儿出院后 30 天内再入院
    h = head.sort_values('入院')
    re30 = {}
    for _, g in h.groupby('科研患者编号'):
        g = g.reset_index(drop=True)
        for i in range(len(g) - 1):
            nxt = g.loc[i + 1, '入院']
            if pd.notna(g.loc[i, '出院']) and nxt - g.loc[i, '出院'] <= pd.Timedelta(days=30):
                re30[g.loc[i, KEY]] = '是（%s 再入院）' % nxt.strftime('%Y-%m-%d')
    a['检索集内30天再入院'] = [re30.get(k, '否') for k in a.index]

    rec_txt = rec.sort_values('t').assign(
        # pandas 3 的字符串类型在 astype(str) 后仍保留缺失值，拼接前须先 fillna
        条目=lambda d: d['t'].dt.strftime('%Y-%m-%d %H:%M').fillna('时间不详') + ' ' + d['手术名称'].fillna('').astype(str)
        + '（术中诊断：' + d['术中诊断'].fillna('').astype(str) + '）：' + d['手术经过'].fillna('').astype(str))
    a['手术记录'] = rec_txt.groupby(KEY)['条目'].apply(lambda s: '\n'.join(s)).reindex(a.index).fillna('')
    for c in ('诊疗经过', '出院情况', '出院医嘱'):
        a[c] = out.groupby(KEY)[c].apply(_join).reindex(a.index).fillna('')
    return a


def final_status():
    """读取终裁结果中的“是否排除”；尚未合并或未定稿者标“裁定未完成”。不读取病因。"""
    if not os.path.exists(MERGED):
        return {}
    fin = pd.read_excel(MERGED, sheet_name='终裁数据集', dtype=object, keep_default_na=False)
    fin[KEY] = fin[KEY].map(norm)
    return {k: ('排除' if e == '是' else '纳入') if s == '已定稿' else '裁定未完成'
            for k, s, e in zip(fin[KEY], fin['终裁状态'], fin['是否排除'].map(norm))}


def build_frame():
    a = auto_frame()
    x = pd.ExcelFile(SRC)
    head = x.parse('病案首页基本信息')
    head[KEY] = head[KEY].map(norm)
    st = final_status()
    df = pd.DataFrame({
        '科研患者编号': head['科研患者编号'],
        KEY: head[KEY],
        '住院号': head['住院号'],
        '第几次住院': head['第几次住院'],
        '入院日期': pd.to_datetime(head['入院日期']).dt.strftime('%Y-%m-%d'),
        '年龄_天': head['年龄（天）'],
        '终裁': head[KEY].map(lambda k: st.get(k, '裁定未完成')),
    })
    auto = ['住院天数', '有手术记录', '术后住院天数', '进食计时起点', '检索集内30天再入院',
            '预估_主诉病程_小时', '提示_影像游离气体', '提示_切口', '提示_再次全麻手术']
    for c in auto:
        df[c] = df[KEY].map(a[c])
    for c in MANUAL_COLS:
        df[c] = pd.Series('', index=df.index, dtype=object)
    for c in HINT_WIDE + TEXT_COLS:
        df[c] = df[KEY].map(a[c])
    df = df.sort_values(['入院日期', KEY]).reset_index(drop=True)
    return df, auto


def carry_over(df, path=None):
    """按就诊编号回填旧文件中已填写的人工列，返回回填条数。"""
    path = path or DST
    if not os.path.exists(path):
        return 0
    old = pd.read_excel(path, sheet_name='提取表', dtype=object, keep_default_na=False)
    if KEY not in old.columns:
        return 0
    old[KEY] = old[KEY].map(norm)
    have = [c for c in MANUAL_COLS if c in old.columns]
    old = old.drop_duplicates(KEY).set_index(KEY)[have]
    n = 0
    for i, k in df[KEY].items():
        if k in old.index:
            for c in have:
                v = old.at[k, c]
                if norm(v) != '':
                    df.at[i, c] = v
                    n += 1
    return n


def _nonneg_or_nr(v):
    if v == 'NR':
        return True
    try:
        return float(v) >= 0
    except ValueError:
        return False


def logic_check(df):
    """按当前已填内容逐例核查。只查已开始提取（填了提取者）且未被排除的病例。"""
    out = []
    grade = {g: i for i, g in enumerate(CLAVIEN)}
    for _, r in df.iterrows():
        v = {c: norm(r[c]) for c in MANUAL_COLS}
        if not v['提取者'] or r['终裁'] == '排除':
            continue

        def add(level, msg):
            out.append({'级别': level, KEY: r[KEY], '住院号': r['住院号'], '入院日期': r['入院日期'], '问题': msg})

        for c, opts in OPTIONS.items():
            if v[c] and v[c] not in opts:
                add('错误', '“%s”填了下拉选项以外的值：%s' % (c, v[c]))
        lack = [c for c in REQUIRED if not v[c]]
        if lack:
            add('错误', '未填：%s（查不到请填“不详”或 NR）' % '、'.join(lack))
        for c in ['主诉病程_小时'] + FEED:
            if v[c] and not _nonneg_or_nr(v[c]):
                add('错误', '“%s”应为非负数或 NR：%s' % (c, v[c]))
        f1, f2 = v[FEED[0]], v[FEED[1]]
        if _nonneg_or_nr(f1) and _nonneg_or_nr(f2) and 'NR' not in (f1, f2) and f1 and f2 and float(f1) > float(f2):
            add('错误', '首次肠内摄入晚于全量肠内营养')
        for c in FEED:
            if v[c] and v[c] != 'NR' and _nonneg_or_nr(v[c]) and pd.notna(r['住院天数']) \
                    and float(v[c]) > float(r['住院天数']) + 1:
                add('错误', '“%s”超过住院天数' % c)
        cd = v['最高Clavien-Dindo分级']
        any_comp = any(v[c] == '是' for c in COMPLICATIONS) or bool(v['其他并发症'])
        if cd == '0' and any_comp:
            add('错误', '有并发症但 Clavien-Dindo 分级为 0')
        if cd in grade and cd != '0' and not any_comp:
            add('错误', 'Clavien-Dindo 分级为 %s，但未勾选任何并发症' % cd)
        if v['非计划再次手术'] == '是':
            if cd in grade and grade[cd] < grade['IIIb']:
                add('错误', '有非计划再次手术，Clavien-Dindo 至少应为 IIIb')
            if not v['再次手术原因']:
                add('错误', '非计划再次手术未写原因')
        if v['非手术再干预'] == '是' and cd in grade and grade[cd] < grade['IIIa']:
            add('错误', '有非手术再干预，Clavien-Dindo 至少应为 IIIa')
        if v['核对者'] and v['核对者'] == v['提取者']:
            add('错误', '核对者与提取者为同一人')
        if v['核对者'] and not v['核对修改说明']:
            add('错误', '已核对但未填“核对修改说明”（无改动写“无”）')
        # 与自动线索不符：提示，逐例确认即可
        if r['有手术记录'] == '否':
            for c in ('切口感染或裂开', '非计划再次手术'):
                if v[c] and v[c] != '不适用':
                    add('提示', '无手术记录，“%s”通常应填“不适用”，请确认' % c)
        if norm(r['提示_切口']) and v['切口感染或裂开'] == '否':
            add('提示', '自动线索有切口问题（%s），但切口感染或裂开填“否”，请确认' % r['提示_切口'])
        if norm(r['提示_再次全麻手术']) and v['非计划再次手术'] == '否':
            add('提示', '首次手术后另有全麻操作（%s），但非计划再次手术填“否”，请确认是否为计划性手术'
                % norm(r['提示_再次全麻手术']))
        if r['提示_影像游离气体'] in ('有', '无') and v['影像游离气体_治疗前'] in ('有', '无') \
                and r['提示_影像游离气体'] != v['影像游离气体_治疗前']:
            add('提示', '影像游离气体与自动线索（%s）不一致，请确认' % r['提示_影像游离气体'])
    return pd.DataFrame(out, columns=['级别', KEY, '住院号', '入院日期', '问题'])


def progress(df):
    need = df[df['终裁'] != '排除']
    ex = need['提取者'].map(norm) != ''
    ck = need['核对者'].map(norm) != ''
    rows = [('需提取例数（不含终裁排除）', len(need)), ('已提取', int(ex.sum())), ('已核对', int(ck.sum())),
            ('终裁判为排除（无需提取）', int((df['终裁'] == '排除').sum())),
            ('裁定未完成', int((df['终裁'] == '裁定未完成').sum()))]
    for c in FEED:
        filled = need[c].map(norm)
        done = filled[filled != '']
        rows.append(('“%s”为 NR 的比例（已填者中）' % c,
                     '%d/%d' % ((done == 'NR').sum(), len(done)) if len(done) else '—'))
    return pd.DataFrame(rows, columns=['项目', '数值'])


def write_workbook(df, auto):
    if os.path.exists(DST):
        bak = DST.replace('.xlsx', '.backup_%s.xlsx' % datetime.now().strftime('%Y%m%d_%H%M%S'))
        shutil.copy2(DST, bak)
        print('已备份旧文件 ->', bak)
    checks = logic_check(df)
    defs = pd.DataFrame(DEFS, columns=['条目', '操作性定义'])
    gray, yellow = PatternFill('solid', fgColor='D9D9D9'), PatternFill('solid', fgColor='FFF2CC')
    dim = PatternFill('solid', fgColor='EDEDED')
    with pd.ExcelWriter(DST, engine='openpyxl') as w:
        defs.to_excel(w, sheet_name='变量定义', index=False)
        df.to_excel(w, sheet_name='提取表', index=False)
        checks.to_excel(w, sheet_name='逻辑核查', index=False)
        progress(df).to_excel(w, sheet_name='进度', index=False)

        ws = w.sheets['提取表']
        cols = list(df.columns)
        for j, name in enumerate(cols, start=1):
            L = get_column_letter(j)
            c = ws.cell(row=1, column=j)
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical='center')
            if name in MANUAL_COLS:
                c.fill = yellow
                ws.column_dimensions[L].width = 12
            elif name in auto or name in HINT_WIDE:
                c.fill = gray
                ws.column_dimensions[L].width = 40 if name in HINT_WIDE else 13
            elif name in TEXT_COLS:
                ws.column_dimensions[L].width = 60
            else:
                ws.column_dimensions[L].width = 12
        for name in ('其他并发症', '再次手术原因', '核对修改说明', '备注'):
            ws.column_dimensions[get_column_letter(cols.index(name) + 1)].width = 24
        ws.freeze_panes = ws.cell(row=2, column=cols.index('终裁') + 2)
        ws.auto_filter.ref = 'A1:%s1' % get_column_letter(len(cols))
        for i, s in enumerate(df['终裁'], start=2):
            if s == '排除':
                for j in range(1, len(cols) + 1):
                    ws.cell(row=i, column=j).fill = dim
        for name, opts in OPTIONS.items():
            L = get_column_letter(cols.index(name) + 1)
            dv = DataValidation(type='list', formula1='"%s"' % ','.join(opts), allow_blank=True)
            ws.add_data_validation(dv)
            dv.add('%s2:%s%d' % (L, L, len(df) + 1))

        dw = w.sheets['变量定义']
        dw.column_dimensions['A'].width = 26
        dw.column_dimensions['B'].width = 110
        for r in range(2, len(defs) + 2):
            dw.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical='top')
        cw = w.sheets['逻辑核查']
        cw.column_dimensions['E'].width = 80
        w.sheets['进度'].column_dimensions['A'].width = 40
    print('已写出 ->', DST)
    return checks


if __name__ == '__main__':
    d, auto_cols = build_frame()
    n = carry_over(d)
    print('汇总 %d 例；从旧文件回填已填写单元格 %d 个' % (len(d), n))
    chk = write_workbook(d, auto_cols)
    print('逻辑核查：错误 %d 条、提示 %d 条' % ((chk['级别'] == '错误').sum(), (chk['级别'] == '提示').sum()))
    print(progress(d).to_string(index=False, header=False))
